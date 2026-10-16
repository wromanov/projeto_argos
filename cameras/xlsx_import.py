"""Leitura, validação e confirmação segura de importações XLSX cadastrais."""

from __future__ import annotations

import copy
import hashlib
import io
import zipfile
import unicodedata
from collections import Counter
from datetime import timedelta

from django.db import transaction
from django.utils import timezone

from .auditoria import calcular_alteracoes, registrar_evento
from .catalogos import CODEC_CHOICES, FARM_CHOICES, RETENCAO_CHOICES
from .forms import CameraCriacaoForm, CameraEdicaoForm
from .integracao.configuration import invalidate_target_config_generation_if_ip_changed
from .models import Camera, CameraImportPreview, EventoAuditoria
from .validadores import normalizar_texto_cadastral

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 5000
PREVIEW_TTL = timedelta(minutes=30)
TEMPLATE_SHEET = "Cameras"
INSTRUCTIONS_SHEET = "Instruções"

# Cabeçalho XLSX -> campo do model. Chaves canônicas sem acentos/minúsculas.
COLUMNS = {
    "operacao": "operacao",
    "patrimonio": "patrimonio",
    "nome": "nome",
    "fabricante": "fabricante",
    "modelo": "modelo",
    "numero de serie": "serial",
    "serial": "serial",
    "mac": "mac",
    "tipo": "tipo",
    "situacao operacional": "situacao_operacional",
    "terminal": "terminal",
    "pavimento": "pavimento",
    "setor": "setor",
    "ambiente": "ambiente",
    "descricao detalhada": "descricao_detalhada",
    "ponto de referencia": "ponto_referencia",
    "ip": "ip",
    "mascara de sub-rede": "mascara_subrede",
    "gateway": "gateway",
    "vlan": "vlan",
    "farm": "farm",
    "retencao esperada (dias)": "retencao_esperada_dias",
    "codec": "codec",
    "observacoes": "observacoes",
}
MODEL_TO_HEADER = {
    "operacao": "OPERAÇÃO", "patrimonio": "PATRIMÔNIO", "nome": "NOME",
    "fabricante": "FABRICANTE", "modelo": "MODELO", "serial": "NÚMERO DE SÉRIE",
    "mac": "MAC", "tipo": "TIPO", "situacao_operacional": "SITUAÇÃO OPERACIONAL",
    "terminal": "TERMINAL", "pavimento": "PAVIMENTO", "setor": "SETOR",
    "ambiente": "AMBIENTE", "descricao_detalhada": "DESCRIÇÃO DETALHADA",
    "ponto_referencia": "PONTO DE REFERÊNCIA", "ip": "IP",
    "mascara_subrede": "MÁSCARA DE SUB-REDE", "gateway": "GATEWAY", "vlan": "VLAN",
    "farm": "FARM", "retencao_esperada_dias": "RETENÇÃO ESPERADA (DIAS)",
    "codec": "CODEC", "observacoes": "OBSERVAÇÕES",
}
IMPORT_FIELDS = tuple(field for field in MODEL_TO_HEADER if field != "operacao")
CLEARABLE_FIELDS = frozenset({
    "patrimonio", "serial", "mac", "tipo", "terminal", "pavimento", "setor",
    "ambiente", "descricao_detalhada", "ponto_referencia", "ip", "mascara_subrede",
    "gateway", "vlan", "farm", "retencao_esperada_dias", "codec", "observacoes",
})
RESTRICTED_HEADERS = {
    "id", "id argos", "status", "monitoramento", "firmware", "firmware version",
    "credencial", "senha", "incidente", "historico", "origem cadastro", "desativada em",
    "desativada por", "motivo desativacao", "criado por", "criado em", "atualizado por",
    "atualizado em", "auditoria",
}


class ImportFileError(ValueError):
    pass


def _key(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).casefold().split())


def _display(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def parse_xlsx(upload):
    if not str(getattr(upload, "name", "")).lower().endswith(".xlsx"):
        raise ImportFileError("O arquivo deve usar a extensão .xlsx.")
    if upload.size > MAX_FILE_BYTES:
        raise ImportFileError("O arquivo excede o limite de 10 MB.")
    raw = upload.read()
    if len(raw) > MAX_FILE_BYTES or not raw.startswith(b"PK"):
        raise ImportFileError("Arquivo XLSX inválido ou acima do limite de 10 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise ImportFileError("O conteúdo descompactado excede o limite de segurança.")
            names = [item.filename.casefold() for item in entries]
            if any(
                name.endswith(("vbaproject.bin", ".bin"))
                or any(part in name for part in ("/activex/", "/embeddings/", "/oleobjects/", "/externallinks/"))
                for name in names
            ):
                raise ImportFileError("Macros e conteúdo binário ativo não são permitidos.")
            if any(
                item.filename.casefold().endswith(".rels")
                and b'targetmode="external"' in archive.read(item).lower()
                for item in entries
            ):
                raise ImportFileError("Links externos não são permitidos.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("Arquivo XLSX inválido.") from exc
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=False, data_only=False, keep_links=True)
    except Exception as exc:
        raise ImportFileError("Não foi possível ler este arquivo XLSX.") from exc
    if wb.sheetnames not in ([TEMPLATE_SHEET], [TEMPLATE_SHEET, INSTRUCTIONS_SHEET]):
        raise ImportFileError(f"A planilha deve conter a aba '{TEMPLATE_SHEET}' e, opcionalmente, '{INSTRUCTIONS_SHEET}'.")
    if getattr(wb, "_external_links", []):
        raise ImportFileError("Links externos não são permitidos.")
    ws = wb[TEMPLATE_SHEET]
    if ws.max_row < 1 or ws.max_row > MAX_DATA_ROWS + 1:
        raise ImportFileError("A planilha deve conter no máximo 5.000 linhas de dados.")
    if ws.max_column > len(MODEL_TO_HEADER):
        raise ImportFileError("A planilha contém colunas excedentes.")
    if any(
        cell.data_type == "f" or cell.hyperlink
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
    ):
        raise ImportFileError("Fórmulas e hiperlinks não são permitidos.")
    headers = [_key(cell.value) for cell in ws[1]]
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise ImportFileError("Há cabeçalhos duplicados.")
    unknown = [h for h in headers if h and h not in COLUMNS]
    if unknown:
        reason = "coluna restrita" if any(h in RESTRICTED_HEADERS for h in unknown) else "coluna desconhecida"
        raise ImportFileError(f"A planilha contém {reason}: {', '.join(unknown)}.")
    mapping = [COLUMNS.get(h) for h in headers]
    mapped_headers = [field for field in mapping if field]
    if len(mapped_headers) != len(set(mapped_headers)):
        raise ImportFileError("Há mais de um cabeçalho para o mesmo campo.")
    if "operacao" not in mapping:
        raise ImportFileError("A coluna OPERAÇÃO é obrigatória.")
    records = []
    for row_number, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in cells):
            continue
        records.append({"line": row_number, "values": {field: _display(value) for field, value in zip(mapping, cells) if field}})
    if not records:
        raise ImportFileError("A planilha não contém linhas de dados.")
    if len(records) > MAX_DATA_ROWS:
        raise ImportFileError("A planilha excede o limite de 5.000 linhas.")
    return raw, records


def _find_target(values):
    patrimonio = values.get("patrimonio") or None
    serial = values.get("serial") or None
    patrimonio = None if patrimonio == "LIMPAR" else patrimonio
    serial = None if serial == "LIMPAR" else serial
    if not patrimonio and not serial:
        return None, "UPDATE exige patrimônio ou número de série."
    by_pat = Camera.objects.filter(patrimonio__iexact=patrimonio).first() if patrimonio else None
    by_serial = Camera.objects.filter(serial__iexact=serial).first() if serial else None
    if by_pat and by_serial and by_pat.pk != by_serial.pk:
        return None, "Patrimônio e número de série apontam para câmeras diferentes."
    target = by_pat or by_serial
    if target:
        # A chave presente, mas não encontrada, só pode complementar o destino encontrado.
        for field, value in (("patrimonio", patrimonio), ("serial", serial)):
            if not value or getattr(target, field) in (None, ""):
                continue
            if getattr(target, field).casefold() != value.casefold():
                return None, f"{field} informado conflita com o cadastro encontrado."
        return target, None
    return None, "UPDATE não encontrou câmera pelo patrimônio ou número de série."


def _form_data(values, instance=None, create=False):
    data = {}
    for field in IMPORT_FIELDS:
        incoming = values.get(field)
        if incoming == "LIMPAR":
            if field not in CLEARABLE_FIELDS:
                raise ValueError(f"LIMPAR não é permitido para {MODEL_TO_HEADER[field]}.")
            incoming = ""
        elif incoming in (None, "") and instance is not None:
            incoming = getattr(instance, field)
        if incoming is None:
            incoming = ""
        data[field] = str(incoming)
    if instance is not None:
        data["origem_cadastro"] = instance.origem_cadastro
        data["situacao_operacional"] = instance.situacao_operacional
    else:
        data["origem_cadastro"] = Camera.OrigemCadastro.IMPORTACAO
        data["situacao_operacional"] = values.get("situacao_operacional", "")
    for field, custom_field, choices in (("farm", "farm_personalizado", FARM_CHOICES), ("codec", "codec_personalizado", CODEC_CHOICES)):
        if data.get(field) and data[field] not in {value for value, _ in choices}:
            data[field] = "PERSONALIZADO"
            data[custom_field] = values[field]
        else:
            data[custom_field] = ""
    retention = data.get("retencao_esperada_dias", "")
    if retention and retention not in {value for value, _ in RETENCAO_CHOICES}:
        data["retencao_esperada_dias"] = "PERSONALIZADO"
        data["retencao_personalizada_dias"] = retention
    else:
        data["retencao_personalizada_dias"] = ""
    if create:
        data["situacao_operacional"] = values.get("situacao_operacional", "")
    return data


def _validated_plan(records):
    plan, errors = [], []
    targets, create_keys = {}, {}
    for record in records:
        line, values = record["line"], dict(record["values"])
        for identifier in ("patrimonio", "serial"):
            if values.get(identifier) not in (None, "", "LIMPAR"):
                values[identifier] = normalizar_texto_cadastral(values[identifier])
        operation = (values.get("operacao") or "").strip().upper()
        row_errors = []
        if operation not in {"CRIAR", "ATUALIZAR"}:
            row_errors.append("OPERAÇÃO deve ser CRIAR ou ATUALIZAR.")
        if operation == "ATUALIZAR" and values.get("situacao_operacional"):
            row_errors.append("Situação operacional não pode ser alterada por UPDATE.")
        instance, form = None, None
        if operation == "ATUALIZAR":
            instance, matching_error = _find_target(values)
            if matching_error:
                row_errors.append(matching_error)
        elif operation == "CRIAR":
            if not values.get("nome") or not values.get("fabricante") or not values.get("modelo"):
                row_errors.append("CREATE exige nome, fabricante e modelo.")
            if values.get("situacao_operacional") not in {"PRODUCAO", "ESTOQUE"}:
                row_errors.append("CREATE exige situação operacional PRODUCAO ou ESTOQUE.")
            if not values.get("patrimonio") and not values.get("serial"):
                row_errors.append("CREATE exige patrimônio ou número de série.")
            for field in ("patrimonio", "serial"):
                value = values.get(field)
                if value:
                    found = Camera.objects.filter(**{f"{field}__iexact": value}).first()
                    if found:
                        row_errors.append(f"CREATE conflita com câmera existente pelo campo {field}.")
                    key = (field, value.casefold())
                    if key in create_keys:
                        row_errors.append(f"Identificador duplicado nas linhas {create_keys[key]} e {line}.")
                    create_keys[key] = line
        original_instance = copy.copy(instance) if instance is not None else None
        if instance:
            previous_line = targets.get(instance.pk)
            if previous_line:
                row_errors.append(f"Esta câmera também é referenciada na linha {previous_line}.")
            targets[instance.pk] = line
        if not row_errors and operation in {"CRIAR", "ATUALIZAR"}:
            try:
                data = _form_data(values, instance=instance, create=operation == "CRIAR")
                form = CameraCriacaoForm(data) if operation == "CRIAR" else CameraEdicaoForm(data, instance=instance)
                if not form.is_valid():
                    row_errors.extend(f"{form.fields[field].label}: {', '.join(messages)}" for field, messages in form.errors.items())
            except (ValueError, TypeError) as exc:
                row_errors.append(str(exc))
        if row_errors:
            errors.extend({"line": line, "message": message} for message in row_errors)
            plan.append({"line": line, "operation": operation, "target_id": instance.pk if instance else None, "errors": row_errors, "values": values})
            continue
        camera = form.save(commit=False)
        if operation == "CRIAR":
            camera.origem_cadastro = Camera.OrigemCadastro.IMPORTACAO
        changes = {}
        if instance:
            changes = calcular_alteracoes(original_instance, camera, IMPORT_FIELDS)
        plan.append({
            "line": line, "operation": operation, "target_id": instance.pk if instance else None,
            "target_label": original_instance.nome if original_instance else "Nova câmera",
            "errors": [], "values": values, "normalized": {field: getattr(camera, field) for field in IMPORT_FIELDS},
            "changes": changes,
        })
    # Identificadores repetidos entre operações diferentes também conflitam
    # quando uma linha UPDATE complementaria a chave que outra linha cria.
    seen_keys = {}
    for row in plan:
        if row["errors"]:
            continue
        for field in ("patrimonio", "serial"):
            value = row.get("normalized", {}).get(field)
            if not value:
                continue
            key = (field, str(value).casefold())
            previous = seen_keys.get(key)
            if previous and previous != row["line"]:
                message = f"Identificador repetido nas linhas {previous} e {row['line']}."
                row["errors"].append(message)
                errors.append({"line": previous, "message": message})
                errors.append({"line": row["line"], "message": message})
            else:
                seen_keys[key] = row["line"]
    # Diferentes linhas não podem atingir a mesma câmera, mesmo com chaves distintas.
    return plan, errors


def preview_xlsx(upload, user):
    raw, records = parse_xlsx(upload)
    file_hash = hashlib.sha256(raw).hexdigest()
    plan, errors = _validated_plan(records)
    operations = {row["operation"] for row in plan}
    if "CRIAR" in operations and not user.has_perm("cameras.add_camera"):
        raise PermissionError("Você não tem permissão para criar câmeras por importação.")
    if "ATUALIZAR" in operations and not user.has_perm("cameras.change_camera"):
        raise PermissionError("Você não tem permissão para atualizar câmeras por importação.")
    conflict_markers = ("conflit", "diferentes", "repetido", "duplicado", "referenciada")
    conflict_lines = {
        error["line"] for error in errors
        if any(marker in error["message"].casefold() for marker in conflict_markers)
    }
    summary = {
        "total": len(plan),
        "create": sum(row["operation"] == "CRIAR" and not row["errors"] for row in plan),
        "update": sum(row["operation"] == "ATUALIZAR" and not row["errors"] for row in plan),
        "invalid": len({error["line"] for error in errors}),
        "conflicts": len(conflict_lines),
        "changes": sum(len(row.get("changes", {})) for row in plan),
        "errors": errors,
        "can_confirm": not errors,
    }
    preview = CameraImportPreview.objects.create(
        usuario=user, file_sha256=file_hash, plano=plan, resumo=summary,
        expira_em=timezone.now() + PREVIEW_TTL,
    )
    return preview


def build_template():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_SHEET
    headers = [MODEL_TO_HEADER[field] for field in MODEL_TO_HEADER]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}1"
    ws.row_dimensions[1].height = 34
    for col, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = max(17, min(31, len(header) + 4))
        ws.cell(1, col).comment = Comment(f"Campo cadastral: {header}. Consulte a aba Instruções.", "ARGOS")
    for row in range(2, MAX_DATA_ROWS + 2):
        for column in (2, 6, 7, 17, 18, 19):
            ws.cell(row, column).number_format = "@"
    op_col = headers.index("OPERAÇÃO") + 1
    operation_validation = DataValidation(type="list", formula1='"CRIAR,ATUALIZAR"', allow_blank=False)
    ws.add_data_validation(operation_validation)
    operation_validation.add(f"{ws.cell(2, op_col).column_letter}2:{ws.cell(MAX_DATA_ROWS + 1, op_col).column_letter}{MAX_DATA_ROWS + 1}")
    situacao_col = headers.index("SITUAÇÃO OPERACIONAL") + 1
    situation_validation = DataValidation(type="list", formula1='"PRODUCAO,ESTOQUE"', allow_blank=True)
    ws.add_data_validation(situation_validation)
    situation_validation.add(f"{ws.cell(2, situacao_col).column_letter}2:{ws.cell(MAX_DATA_ROWS + 1, situacao_col).column_letter}{MAX_DATA_ROWS + 1}")
    tipo_col = headers.index("TIPO") + 1
    type_choices = ",".join(Camera.TipoCamera.values)
    type_validation = DataValidation(type="list", formula1=f'"{type_choices}"', allow_blank=True)
    ws.add_data_validation(type_validation)
    type_validation.add(f"{ws.cell(2, tipo_col).column_letter}2:{ws.cell(MAX_DATA_ROWS + 1, tipo_col).column_letter}{MAX_DATA_ROWS + 1}")
    instructions = wb.create_sheet(INSTRUCTIONS_SHEET)
    instructions.append(["Importação cadastral de câmeras — template v1"])
    instructions.append(["Preencha a aba Cameras. Não altere nem acrescente cabeçalhos."])
    instructions.append(["Limites: até 5.000 linhas de dados e arquivo de até 10 MB."])
    instructions.append(["OPERAÇÃO: CRIAR ou ATUALIZAR. UPDATE usa somente patrimônio e/ou número de série."])
    instructions.append(["CREATE exige nome, fabricante, modelo, situação (PRODUCAO/ESTOQUE) e patrimônio ou número de série."])
    instructions.append(["Em UPDATE, célula vazia preserva o valor. Use LIMPAR apenas em campos opcionais que podem ser limpos."])
    instructions.append(["UPDATE não pode alterar situação operacional. Não inclua ID, monitoramento, firmware, credenciais, incidentes ou histórico."])
    instructions.append(["COLUNA", "TIPO E REGRAS"])
    specifications = {
        "operacao": "Texto obrigatório: CRIAR ou ATUALIZAR.",
        "patrimonio": "Texto (máx. 50); chave para matching; CREATE exige patrimônio ou série; LIMPAR permitido em UPDATE.",
        "nome": "Texto obrigatório em CREATE (máx. 150); LIMPAR proibido.",
        "fabricante": "Texto obrigatório em CREATE (máx. 100); LIMPAR proibido.",
        "modelo": "Texto obrigatório em CREATE (máx. 100); LIMPAR proibido.",
        "serial": "Texto (máx. 100); chave para matching; CREATE exige patrimônio ou série; LIMPAR permitido em UPDATE.",
        "mac": "Texto no formato MAC; deve ser único; LIMPAR permitido em UPDATE.",
        "tipo": "Choice: BULLET, DOME, SPEED_DOME ou FISHEYE; LIMPAR permitido em UPDATE.",
        "situacao_operacional": "CREATE obrigatório: PRODUCAO ou ESTOQUE. Deve ficar vazio em UPDATE.",
        "ip": "Endereço IPv4; LIMPAR permitido em UPDATE.",
        "mascara_subrede": "Endereço IPv4; LIMPAR permitido em UPDATE.",
        "gateway": "Endereço IPv4; LIMPAR permitido em UPDATE.",
        "vlan": "Inteiro entre 1 e 4094; LIMPAR permitido em UPDATE.",
        "retencao_esperada_dias": "Inteiro maior ou igual a 1; LIMPAR permitido em UPDATE.",
        "farm": "Texto do catálogo ou valor personalizado; LIMPAR permitido em UPDATE.",
        "codec": "Texto do catálogo ou valor personalizado; LIMPAR permitido em UPDATE.",
    }
    optional_text = {"mac", "tipo", "ip", "mascara_subrede", "gateway", "vlan", "retencao_esperada_dias", "farm", "codec"}
    for field, header in MODEL_TO_HEADER.items():
        if field == "operacao" or field in specifications:
            detail = specifications.get(field)
        elif field in optional_text:
            detail = "Campo opcional; célula vazia preserva em UPDATE; LIMPAR permitido em UPDATE."
        else:
            detail = "Texto opcional; célula vazia preserva em UPDATE; LIMPAR permitido em UPDATE."
        instructions.append([header, detail])
    instructions.append(["Choices de farm: " + ", ".join(value for value, _ in FARM_CHOICES if value and value != "PERSONALIZADO")])
    instructions.append(["Choices de codec: " + ", ".join(value for value, _ in CODEC_CHOICES if value and value != "PERSONALIZADO")])
    instructions.append(["Choices de retenção (dias): " + ", ".join(value for value, _ in RETENCAO_CHOICES if value and value != "PERSONALIZADO")])
    instructions.append(["Sem fórmulas, macros, hiperlinks ou conteúdo externo. Uma linha inválida bloqueia o lote inteiro."])
    instructions.column_dimensions["A"].width = 34
    instructions.column_dimensions["B"].width = 105
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _revalidate_preview(preview, *, lock=False):
    if lock:
        target_ids = sorted({row.get("target_id") for row in preview.plano if row.get("target_id")})
        list(Camera.objects.select_for_update().filter(pk__in=target_ids).order_by("pk"))
    plan, errors = _validated_plan(preview.plano)
    if errors:
        raise ValueError("Os dados ou identificadores mudaram desde a prévia. Valide o arquivo novamente.")
    old_signature = [(r["line"], r["operation"], r.get("target_id"), r.get("normalized"), r.get("changes")) for r in preview.plano]
    new_signature = [(r["line"], r["operation"], r.get("target_id"), r.get("normalized"), r.get("changes")) for r in plan]
    if old_signature != new_signature:
        raise ValueError("O cadastro mudou desde a prévia. Valide o arquivo novamente.")
    return plan


@transaction.atomic
def confirm_preview(preview, user, request=None):
    preview = CameraImportPreview.objects.select_for_update().get(pk=preview.pk)
    if preview.usuario_id != user.pk:
        raise PermissionError("Esta prévia pertence a outro usuário.")
    if preview.consumido_em:
        raise ValueError("Esta prévia já foi confirmada.")
    if preview.expira_em <= timezone.now():
        raise TimeoutError("A prévia expirou. Envie o arquivo novamente.")
    plan = _revalidate_preview(preview, lock=True)
    batch_id = str(preview.pk)
    results = Counter()
    for row in plan:
        if row["operation"] == "CRIAR":
            form = CameraCriacaoForm(_form_data(row["values"], create=True))
            if not form.is_valid():
                raise ValueError(f"Linha {row['line']}: os dados não são mais válidos.")
            camera = form.save(commit=False)
            camera.origem_cadastro = Camera.OrigemCadastro.IMPORTACAO
            camera.criado_por = camera.atualizado_por = user
            camera.save()
            changes = {key: {"anterior": None, "atual": getattr(camera, key)} for key in IMPORT_FIELDS}
            results["created"] += 1
            action = "CRIAR"
            previous_ip = None
        else:
            camera = Camera.objects.select_for_update().get(pk=row["target_id"])
            before = Camera.objects.get(pk=camera.pk)
            data = _form_data(row["values"], instance=before)
            form = CameraEdicaoForm(data, instance=camera)
            if not form.is_valid():
                raise ValueError(f"Linha {row['line']}: os dados não são mais válidos.")
            camera = form.save(commit=False)
            previous_ip = before.ip
            changes = calcular_alteracoes(before, camera, IMPORT_FIELDS)
            camera.atualizado_por = user
            camera.save()
            invalidate_target_config_generation_if_ip_changed(camera_id=camera.pk, previous_ip=previous_ip, current_ip=camera.ip)
            results["updated"] += 1
            action = "ATUALIZAR"
        registrar_evento(
            codigo="camera.importada", resumo=f"Câmera {action.lower()} por importação XLSX",
            categoria=EventoAuditoria.Categoria.CAMERA, usuario=user, camera=camera,
            entidade=camera, alteracoes=changes,
            metadados={"lote_id": batch_id, "linha": row["line"], "operacao": action},
        )
    registrar_evento(
        codigo="camera.importacao.concluida", resumo="Importação XLSX concluída",
        categoria=EventoAuditoria.Categoria.SISTEMA, usuario=user,
        metadados={"lote_id": batch_id, "total": len(plan), "criadas": results["created"], "atualizadas": results["updated"], "resultado": "SUCESSO"},
    )
    preview.consumido_em = timezone.now()
    preview.save(update_fields=("consumido_em",))
    return {"lote_id": batch_id, "created": results["created"], "updated": results["updated"], "total": len(plan)}


def record_failed_confirmation(preview, user):
    """Registra fora da transação revertida que a confirmação não foi aplicada."""
    registrar_evento(
        codigo="camera.importacao.falhou", resumo="Confirmação de importação XLSX não aplicada",
        categoria=EventoAuditoria.Categoria.SISTEMA, usuario=user,
        metadados={"lote_id": str(preview.pk), "resultado": "ROLLBACK"},
    )
